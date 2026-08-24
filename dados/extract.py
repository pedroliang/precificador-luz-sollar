import openpyxl, re, json, math
from openpyxl.utils import get_column_letter

WB = '/home/claude/data/book.xlsx'
wb = openpyxl.load_workbook(WB, data_only=False)
wbv = openpyxl.load_workbook(WB, data_only=True)

def num(v):
    if v is None: return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)): return None
        return round(float(v), 6)
    return None

def parsebrl(v):
    if v is None: return None
    if isinstance(v, (int, float)): return num(v)
    s = str(v).strip()
    if not s or s == '-': return None
    s = s.replace('R$', '').replace(' ', '').replace('\xa0', '')
    s = s.replace('.', '').replace(',', '.')
    try: return round(float(s), 6)
    except ValueError: return None

def rnum(ws, wsv, r, c):
    """value of a cell, preferring literal, else cached computed"""
    v = ws.cell(r, c).value
    if isinstance(v, str) and v.startswith('='):
        return num(wsv.cell(r, c).value)
    return parsebrl(v)

def override(ws, r, c):
    """literal (manually typed) value in a normally-formula column, else None"""
    v = ws.cell(r, c).value
    if v is None: return None
    if isinstance(v, str) and v.startswith('='): return None
    return parsebrl(v)

# ---------- Custo Produtos ----------
ws = wb['Custo Produtos']; wsv = wbv['Custo Produtos']
produtos = {}
for r in range(2, ws.max_row + 1):
    sku = ws.cell(r, 1).value
    desc = ws.cell(r, 2).value
    custo = rnum(ws, wsv, r, 3)
    if sku is None: continue
    sku = str(sku).strip()
    if sku.endswith('.0'): sku = sku[:-2]
    if not sku: continue
    produtos[sku] = [desc if desc else '', custo]

# ---------- store configs ----------
MELI = ['GERAL MELI LUZ SOLLAR', 'GERAL MELI E-STECK', 'GERAL MELI E-BRILIA',
        'GERAL MELI E-LUMIN', 'GERAL MELI APOLLO', 'GERAL MELI CENTRAL LUX ']

def norm_sku(v):
    if v is None: return ''
    if isinstance(v, float) and v == int(v): return str(int(v))
    return str(v).strip()

def parse_custo(cell, row_of, sheetrow):
    """returns rule"""
    v = cell.value
    if v is None: return None
    if isinstance(v, str) and v.startswith('='):
        s = v.replace(' ', '')
        if s.upper().startswith('=VLOOKUP'):
            return ['L']            # lookup by own sku
        m = re.match(r'^=C(\d+)\*([\d.]+)$', s)
        if m:
            return ['M', int(m.group(1)), float(m.group(2))]
        m = re.match(r'^=C(\d+)\+C(\d+)$', s)
        if m:
            return ['S', int(m.group(1)), int(m.group(2))]
        return None
    return ['F', num(v)]

stores = []

def extract_meli(name):
    ws = wb[name]; wsv = wbv[name]
    defaults = {k: num(ws.cell(2, c).value) for k, c in
                [('nf', 7), ('ads', 8), ('lucro', 9), ('dev', 10), ('cls', 11), ('prm', 12)]}
    rows = []
    idx = {}
    for r in range(3, ws.max_row + 1):
        sku = norm_sku(ws.cell(r, 1).value)
        prod = ws.cell(r, 2).value
        rule = parse_custo(ws.cell(r, 3), None, r)
        if prod is None and rule is None: continue
        idx[r] = len(rows)
        rates = {}
        for k, c in [('nf', 7), ('ads', 8), ('lucro', 9), ('dev', 10), ('cls', 11), ('prm', 12)]:
            cv = ws.cell(r, c).value
            if isinstance(cv, str) and cv.startswith('=$'):
                pass
            else:
                n = num(cv)
                if n is not None and n != defaults[k]:
                    rates[k] = n
        rows.append({
            'sku': sku,
            'nome': prod or '',
            'cr': rule,
            'cv': rnum(ws, wsv, r, 3),
            'desc': num(ws.cell(r, 4).value) or 0,
            'eov': override(ws, r, 5),
            'frete': num(ws.cell(r, 13).value) or 0,
            'rates': rates,
            'promoC': rnum(ws, wsv, r, 18),
            'rebC': rnum(ws, wsv, r, 20),
            'promoP': rnum(ws, wsv, r, 27),
            'rebP': rnum(ws, wsv, r, 29),
            'ov': {k: v for k, v in {
                'cls_min': override(ws, r, 15), 'cls_ven': override(ws, r, 16), 'cls_ori': override(ws, r, 17),
                'prm_min': override(ws, r, 24), 'prm_ven': override(ws, r, 25), 'prm_ori': override(ws, r, 26),
            }.items() if v is not None},
            'mkt': [rnum(ws, wsv, r, c) for c in range(33, 39)],
            'obs': ws.cell(r, 39).value or '',
            '_r': r,
        })
    # remap formula row refs -> array indices
    for row in rows:
        cr = row['cr']
        if cr and cr[0] == 'M':
            cr[1] = idx.get(cr[1], -1)
        elif cr and cr[0] == 'S':
            cr[1] = idx.get(cr[1], -1); cr[2] = idx.get(cr[2], -1)
        row['r'] = row.pop('_r')
    return {
        'id': name,
        'nome': name.replace('GERAL MELI ', '').replace('GERAL ', '').strip().title(),
        'plat': 'meli',
        'defaults': defaults,
        'fixos': [{'key': 'frete', 'label': 'Frete/Fixo'}],
        'canais': [{'key': 'cls', 'label': 'Clássico'}, {'key': 'prm', 'label': 'Premium'}],
        'rows': rows,
    }

def extract_generic(name, cols, fixos, ratekeys, comkey):
    """cols: dict of column numbers"""
    ws = wb[name]; wsv = wbv[name]
    defaults = {}
    for k, c in ratekeys.items():
        defaults[k] = num(ws.cell(2, c).value) or 0
    rows = []
    idx = {}
    for r in range(3, ws.max_row + 1):
        sku = norm_sku(ws.cell(r, 1).value)
        prod = ws.cell(r, 2).value
        rule = parse_custo(ws.cell(r, 3), None, r)
        if prod is None and rule is None: continue
        idx[r] = len(rows)
        rates = {}
        for k, c in ratekeys.items():
            cv = ws.cell(r, c).value
            if isinstance(cv, str) and cv.startswith('=$'):
                continue
            n = num(cv)
            if n is not None and n != defaults.get(k):
                rates[k] = n
        n = num(ws.cell(r, cols['com']).value)
        rates[comkey] = n if n is not None else 0
        d = {
            'sku': sku,
            'nome': prod or '',
            'cr': rule,
            'cv': rnum(ws, wsv, r, 3),
            'desc': num(ws.cell(r, 4).value) or 0,
            'eov': override(ws, r, 5),
            'rates': rates,
            'promoC': rnum(ws, wsv, r, cols['promo']),
            'rebC': None, 'promoP': None, 'rebP': None,
            'ov': {k: v for k, v in {
                comkey + '_min': override(ws, r, cols['min']),
                comkey + '_ven': override(ws, r, cols['ven']),
                comkey + '_ori': override(ws, r, cols['ori']),
            }.items() if v is not None},
            'mkt': [None] * 6,
            'obs': '',
            '_r': r,
        }
        for f in fixos:
            d[f['key']] = num(ws.cell(r, f['col']).value) or 0
        rows.append(d)
    for row in rows:
        cr = row['cr']
        if cr and cr[0] == 'M':
            cr[1] = idx.get(cr[1], -1)
        elif cr and cr[0] == 'S':
            cr[1] = idx.get(cr[1], -1); cr[2] = idx.get(cr[2], -1)
        row['r'] = row.pop('_r')
    return {
        'id': name,
        'nome': name.replace('GERAL ', '').strip().title(),
        'plat': 'other',
        'defaults': defaults,
        'fixos': [{'key': f['key'], 'label': f['label']} for f in fixos],
        'canais': [{'key': comkey, 'label': 'Padrão'}],
        'rows': rows,
    }

for name in MELI:
    stores.append(extract_meli(name))

stores.append(extract_generic(
    'GERAL SHOPEE',
    {'com': 11, 'promo': 17, 'min': 14, 'ven': 15, 'ori': 16},
    [{'key': 'fixo', 'label': 'Taxa Fixa (R$)', 'col': 12}],
    {'nf': 7, 'ads': 8, 'lucro': 9, 'dev': 10},
    'com'))

stores.append(extract_generic(
    'GERAL TIKTOK',
    {'com': 11, 'promo': 18, 'min': 15, 'ven': 16, 'ori': 17},
    [{'key': 'taxa', 'label': 'Taxa (R$)', 'col': 12}, {'key': 'frete', 'label': 'Frete (R$)', 'col': 13}],
    {'nf': 7, 'ads': 8, 'lucro': 9, 'dev': 10},
    'com'))

out = {'produtos': produtos, 'stores': stores}
with open('/home/claude/data/dataset.json', 'w') as f:
    json.dump(out, f, ensure_ascii=False, separators=(',', ':'))

for s in stores:
    print(s['id'], len(s['rows']), s['defaults'])
print('produtos', len(produtos))
import os
print('size', os.path.getsize('/home/claude/data/dataset.json'))
